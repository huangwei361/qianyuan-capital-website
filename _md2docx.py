"""Markdown → Word 转换（处理标题/表格/列表/引用/加粗）"""
import re, sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r'C:\Users\64549\.minimax\金融财富公司成立\乾元量化·多因子选股策略·V1.md'
OUT = r'C:\Users\64549\.minimax\金融财富公司成立\乾元量化·多因子选股策略·V1.docx'

with open(SRC, 'r', encoding='utf-8') as f:
    md = f.read()

# 设置中文字体
def set_cn_font(run, size_pt=11, bold=False, color=None, font_name='Microsoft YaHei'):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    r = run._element
    rPr = r.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        r.insert(0, rPr)
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font_name)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    if color:
        run.font.color.rgb = RGBColor(*color)

doc = Document()

# 全局样式
styles = doc.styles
norm = styles['Normal']
norm.font.name = 'Microsoft YaHei'
norm.font.size = Pt(11)
rPr = norm.element.get_or_add_rPr()
rFonts = rPr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rPr.append(rFonts)
rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

# 页面边距
for section in doc.sections:
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

# 内联格式解析：**bold** *italic* `code`
def parse_inline(paragraph, text):
    """处理行内格式：先识别 pattern，按顺序添加 runs"""
    pattern = r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            set_cn_font(run, bold=True)
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            set_cn_font(run)
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            set_cn_font(run, font_name='Consolas', color=(120, 60, 0))
        else:
            run = paragraph.add_run(part)
            set_cn_font(run)

def add_heading(level, text):
    h = doc.add_heading('', level=level)
    parse_inline(h, text)
    for run in h.runs:
        if level == 1:
            set_cn_font(run, size_pt=18, bold=True, color=(6, 24, 42))
        elif level == 2:
            set_cn_font(run, size_pt=15, bold=True, color=(6, 24, 42))
        elif level == 3:
            set_cn_font(run, size_pt=13, bold=True, color=(184, 134, 11))

def add_quote(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.right_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '18')
    left.set(qn('w:space'), '10')
    left.set(qn('w:color'), 'D4AF37')
    pBdr.append(left)
    pPr.append(pBdr)
    parse_inline(p, text.lstrip('> ').strip())
    for run in p.runs:
        set_cn_font(run, size_pt=11, color=(102, 102, 102))

def add_list_item(text, ordered=False):
    style = 'List Number' if ordered else 'List Bullet'
    p = doc.add_paragraph(style=style)
    parse_inline(p, text)
    for run in p.runs:
        set_cn_font(run)

def add_paragraph(text):
    p = doc.add_paragraph()
    parse_inline(p, text)
    for run in p.runs:
        set_cn_font(run)

def add_table(rows):
    """rows: 列表，每行是 list of cell strings"""
    if not rows:
        return
    cols = len(rows[0])
    tbl = doc.add_table(rows=len(rows), cols=cols)
    tbl.style = 'Light Grid Accent 1'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = tbl.rows[i].cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            parse_inline(p, cell_text.strip())
            for run in p.runs:
                if i == 0:
                    set_cn_font(run, size_pt=10, bold=True, color=(255, 255, 255))
                else:
                    set_cn_font(run, size_pt=10)
            # 表头底色
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), '06182A')
                tcPr.append(shd)

def add_separator():
    p = doc.add_paragraph()
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'D4AF37')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ============ 解析 ============
lines = md.split('\n')
i = 0
n = len(lines)
while i < n:
    line = lines[i]
    stripped = line.rstrip()

    # 空行跳过
    if not stripped:
        i += 1
        continue

    # 分隔线
    if stripped == '---':
        add_separator()
        i += 1
        continue

    # 标题
    if stripped.startswith('#### '):
        add_heading(4, stripped[5:])
        i += 1
        continue
    if stripped.startswith('### '):
        add_heading(3, stripped[4:])
        i += 1
        continue
    if stripped.startswith('## '):
        add_heading(2, stripped[3:])
        i += 1
        continue
    if stripped.startswith('# '):
        add_heading(1, stripped[2:])
        i += 1
        continue

    # 引用
    if stripped.startswith('> '):
        add_quote(stripped)
        i += 1
        continue

    # 表格（连续多行以 | 开头）
    if stripped.startswith('|'):
        tbl_lines = []
        while i < n and lines[i].rstrip().startswith('|'):
            tbl_lines.append(lines[i].rstrip())
            i += 1
        # 解析表格行：去掉首尾 |，按 | 分割
        rows = []
        for tl in tbl_lines:
            # 跳过分隔行 |---|---|
            if re.match(r'^\|[\s\-:|]+\|?$', tl):
                continue
            cells = [c.strip() for c in tl.strip('|').split('|')]
            rows.append(cells)
        if rows:
            add_table(rows)
        continue

    # 列表
    if re.match(r'^\s*[-*]\s+', stripped):
        text = re.sub(r'^\s*[-*]\s+', '', stripped)
        add_list_item(text, ordered=False)
        i += 1
        continue
    if re.match(r'^\s*\d+\.\s+', stripped):
        text = re.sub(r'^\s*\d+\.\s+', '', stripped)
        add_list_item(text, ordered=True)
        i += 1
        continue

    # 普通段落
    add_paragraph(stripped)
    i += 1

doc.save(OUT)
import os
print(f'OK: {OUT}')
print(f'Size: {os.path.getsize(OUT)} bytes')
print(f'Paragraphs: {len(doc.paragraphs)}')